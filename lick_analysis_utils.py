# lick_analysis_utils.py
# Análisis ligero por sesión + actualización de libro maestro por día

import csv
import pandas as pd
import numpy as np
from pathlib import Path
from openpyxl import load_workbook


def _normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Normaliza nombres de columnas y mapea aliases comunes a 'delta_ms'."""
    if df is None or df.empty:
        return df

    df = df.copy()
    df.columns = [str(c).strip().replace("\ufeff", "") for c in df.columns]

    alias_map = {
        "delta_ms": "delta_ms",
        "delta": "delta_ms",
        "dt_ms": "delta_ms",
        "ili": "delta_ms",
        "ili_ms": "delta_ms",
        "interlickinterval_ms": "delta_ms",
        "inter_lick_interval_ms": "delta_ms",
    }

    lowered = {c.lower().replace(" ", "").replace("-", "_"): c for c in df.columns}
    for key, target in alias_map.items():
        k = key.lower().replace(" ", "").replace("-", "_")
        if k in lowered and target not in df.columns:
            df = df.rename(columns={lowered[k]: target})
            break

    return df


def _coerce_single_numeric_to_delta(df: pd.DataFrame) -> pd.DataFrame:
    """Si no existe delta_ms pero hay una sola columna numérica utilizable, renombrarla."""
    if df is None or df.empty or "delta_ms" in df.columns:
        return df

    tmp = df.copy()
    numeric_candidates = []
    for col in tmp.columns:
        s = pd.to_numeric(tmp[col], errors="coerce")
        n_valid = int(s.notna().sum())
        if n_valid >= 3:
            numeric_candidates.append((col, n_valid))

    if len(numeric_candidates) == 1:
        only_col = numeric_candidates[0][0]
        tmp = tmp.rename(columns={only_col: "delta_ms"})
        return tmp

    return df


# -----------------------------
# Parámetros globales
# -----------------------------
THRESHOLD_MS = 250
MIN_ILI_MS = 75
ICI_MIN = 500
LONG_MIN = 1000


# =========================================================
# Microestructura núcleo (misma lógica que Enero_2026)
# =========================================================

def burst_metrics_from_ilis(ilis_ms):
    ilis = [float(x) for x in ilis_ms if float(x) >= MIN_ILI_MS]

    bursts_dur = []
    bursts_licks = []
    ibis = []
    icis = []
    icis_long = []

    run = []

    for x in ilis:
        if x <= THRESHOLD_MS:
            run.append(x)
        else:
            if len(run) >= 2:
                bursts_dur.append(sum(run))
                bursts_licks.append(len(run) + 1)

                # IBI = cualquier ILI > 250 que rompe ráfaga válida
                ibis.append(x)

                if x > LONG_MIN:
                    icis_long.append(x)
                    icis.append(x)
                elif x > ICI_MIN:
                    icis.append(x)

            run = []

    if len(run) >= 2:
        bursts_dur.append(sum(run))
        bursts_licks.append(len(run) + 1)

    # Clusters <= 500 ms
    clusters = []
    c_run = []
    for x in ilis:
        if x <= ICI_MIN:
            c_run.append(x)
        else:
            if c_run:
                clusters.append(sum(c_run))
                c_run = []
    if c_run:
        clusters.append(sum(c_run))

    return {
        "n_bursts": len(bursts_dur),
        "avg_burst_duration_ms": np.mean(bursts_dur) if bursts_dur else 0,
        "avg_licks_per_burst": np.mean(bursts_licks) if bursts_licks else 0,
        "avg_ibi_ms": np.mean(ibis) if ibis else 0,
        "avg_ici_ms": np.mean(icis) if icis else 0,
        "avg_ici_long_ms": np.mean(icis_long) if icis_long else 0,
        "avg_cluster_size_ms": np.mean(clusters) if clusters else 0,
    }


# =========================================================
# Licks por minuto
# =========================================================

def licks_per_minute(ilis):
    if not ilis:
        return np.zeros(10)

    time_end = np.cumsum(ilis)
    minutes = (time_end // 60000).astype(int)

    arr = np.zeros(10)
    for m in minutes:
        if 0 <= m < 10:
            arr[int(m)] += 1

    return arr


def _read_lick_csv_robust(csv_path):
    """Lee CSVs del lickómetro (incluyendo CH1 con metadata arriba y tabla real más abajo)."""
    path = Path(csv_path)

    # Intento 0 (prioritario): detectar header real después de metadata, ej:
    # indice_lick,t_rel_ms,delta_ms
    try:
        with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as f:
            lines = f.readlines()

        header_idx = None
        sep = ","
        for i, line in enumerate(lines[:200]):
            txt = line.strip().lower().replace(" ", "")
            if ("delta_ms" in txt) and (("indice_lick" in txt) or ("t_rel_ms" in txt) or ("t_rel" in txt)):
                header_idx = i
                if "\t" in line:
                    sep = "\t"
                elif ";" in line:
                    sep = ";"
                else:
                    sep = ","
                break

        if header_idx is not None:
            df = pd.read_csv(
                path,
                skiprows=header_idx,
                sep=sep,
                engine="python",
                on_bad_lines="skip",
            )
            df = _normalize_columns(df)
            df = _coerce_single_numeric_to_delta(df)
            if "delta_ms" in df.columns:
                return df
    except Exception:
        pass

    # Intento 1: lectura estándar
    try:
        df = pd.read_csv(path)
        df = _normalize_columns(df)
        df = _coerce_single_numeric_to_delta(df)
        if "delta_ms" in df.columns:
            return df
    except pd.errors.ParserError:
        pass
    except Exception:
        pass

    # Intento 2: detectar delimitador (coma, ; o tab)
    try:
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            sample = f.read(4096)
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        sep = dialect.delimiter
    except Exception:
        sep = ","

    try:
        df = pd.read_csv(path, sep=sep, engine="python", on_bad_lines="skip")
        df = _normalize_columns(df)
        df = _coerce_single_numeric_to_delta(df)
        if "delta_ms" in df.columns:
            return df
    except Exception:
        pass

    # Intento 3: forzar separación automática con engine python
    try:
        df = pd.read_csv(path, sep=None, engine="python", on_bad_lines="skip")
        df = _normalize_columns(df)
        df = _coerce_single_numeric_to_delta(df)
        if "delta_ms" in df.columns:
            return df
    except Exception:
        pass

    # Intento 4: si el archivo es una sola columna con texto tipo CSV, expandir
    try:
        raw = pd.read_csv(path, header=None, dtype=str, engine="python", on_bad_lines="skip")
        if raw.shape[1] == 1:
            col0 = raw.iloc[:, 0].astype(str)
            header_idx = None
            for i, val in enumerate(col0.tolist()[:200]):
                low = str(val).lower().replace(" ", "")
                if "delta_ms" in low:
                    header_idx = i
                    break

            if header_idx is not None:
                header_text = str(col0.iloc[header_idx])
                sep_guess = ","
                if "\t" in header_text:
                    sep_guess = "\t"
                elif ";" in header_text:
                    sep_guess = ";"

                split_rows = col0.iloc[header_idx:]
                # Normalizar delimitador a coma para el split
                split_rows = split_rows.str.replace("\t", ",", regex=False)
                if sep_guess == ";":
                    split_rows = split_rows.str.replace(";", ",", regex=False)
                split_rows = split_rows.str.split(",", expand=True)
                split_rows.columns = split_rows.iloc[0].tolist()
                df = split_rows.iloc[1:].reset_index(drop=True)
                df = _normalize_columns(df)
                df = _coerce_single_numeric_to_delta(df)
                if "delta_ms" in df.columns:
                    return df
    except Exception:
        pass

    # Intento 5: extracción manual desde líneas sucias (último recurso)
    try:
        values = []
        with open(path, "r", encoding="utf-8-sig", errors="ignore") as f:
            for line in f:
                txt = line.strip()
                if not txt:
                    continue

                low = txt.lower()
                if any(k in low for k in ["delta_ms", "rat_id", "timestamp", "lick", "session", "indice_lick", "t_rel_ms"]):
                    continue

                parts = txt.replace(";", ",").replace("\t", ",").split(",")
                parsed_any = False
                for p in parts:
                    p = p.strip()
                    if not p:
                        continue
                    try:
                        v = float(p)
                        if np.isfinite(v) and v >= 0:
                            values.append(v)
                            parsed_any = True
                            break  # primer numérico por línea
                    except Exception:
                        continue

                if (not parsed_any) and ("," not in txt):
                    try:
                        v = float(txt)
                        if np.isfinite(v) and v >= 0:
                            values.append(v)
                    except Exception:
                        pass

        if len(values) >= 3:
            return pd.DataFrame({"delta_ms": values})
    except Exception:
        pass

    raise ValueError(
        f"No se pudo leer el CSV: {path}. Verifica delimitador, encabezados y que exista la columna 'delta_ms'."
    )


# =========================================================
# Función principal llamada por el lickómetro
# =========================================================

def analyze_and_append(csv_path, day_label, output_folder):

    df = _read_lick_csv_robust(csv_path)
    df = _normalize_columns(df)

    # Asegúrate que tu CSV tenga columna delta_ms
    if "delta_ms" not in df.columns:
        raise KeyError(
            f"La columna 'delta_ms' no existe en {csv_path}. Columnas detectadas: {list(df.columns)}"
        )

    # Limpieza robusta por si hay texto, espacios o filas corruptas
    df["delta_ms"] = pd.to_numeric(df["delta_ms"], errors="coerce")
    df = df.dropna(subset=["delta_ms"]).copy()
    # quitar primer delta=0 (muy común en CH1) y cualquier valor no positivo
    df = df[df["delta_ms"] > 0].copy()

    ilis = df["delta_ms"].astype(float).tolist()

    if len(ilis) == 0:
        raise ValueError(f"No hay valores válidos en 'delta_ms' para analizar en: {csv_path}")

    # ==============================
    # MÉTRICAS COMPLETAS (10 min)
    # ==============================

    metrics_full = burst_metrics_from_ilis(ilis)
    lpm = licks_per_minute(ilis)

    rat_id = df.iloc[0]["rat_id"] if "rat_id" in df.columns else Path(csv_path).stem

    row_full = {
        "rat_id": rat_id,
        "session_duration_min": round(sum(ilis) / 60000, 2),
        "licks_total": len(ilis),
        "licks_0_3min": sum(lpm[0:3]),
        "licks_0_5min": sum(lpm[0:5]),
        **metrics_full
    }

    # ==============================
    # MÉTRICAS PRIMER MINUTO
    # ==============================

    time_end = np.cumsum(ilis)
    first_mask = time_end <= 60000
    ilis_first = list(np.array(ilis)[first_mask])

    metrics_first = burst_metrics_from_ilis(ilis_first)

    row_first = {
        "rat_id": rat_id,
        "licks_first_min": len(ilis_first),
        **metrics_first
    }

    # ==============================
    # Guardado
    # ==============================

    master_path = Path(output_folder) / f"{day_label}_resultados.xlsx"

    if master_path.exists():

        wb = load_workbook(master_path)

        # ---- Resumen completo
        ws_full = wb["Resumen"]
        ws_full.append(list(row_full.values()))

        # ---- Primer minuto
        if "Primer_minuto" not in wb.sheetnames:
            ws_first = wb.create_sheet("Primer_minuto")
            ws_first.append(list(row_first.keys()))
        else:
            ws_first = wb["Primer_minuto"]

        ws_first.append(list(row_first.values()))

        wb.save(master_path)

    else:
        with pd.ExcelWriter(master_path, engine="openpyxl") as writer:
            pd.DataFrame([row_full]).to_excel(
                writer, sheet_name="Resumen", index=False
            )
            pd.DataFrame([row_first]).to_excel(
                writer, sheet_name="Primer_minuto", index=False
            )

    print(f"✅ Actualizado {master_path}")